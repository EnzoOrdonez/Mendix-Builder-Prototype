"""Script para generar el fixture sample_form.xlsx.

Uso:
    python scripts/create_sample_xlsx.py
"""

from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl no instalado. Ejecuta: pip install openpyxl")
    raise SystemExit(1)


def create_sample_form() -> None:
    """Crea el archivo fixtures/sample_form.xlsx con datos de ejemplo."""
    wb = openpyxl.Workbook()

    # ─── Hoja 1: OrdenCompra ────────────────────────────────
    ws1 = wb.active
    ws1.title = "OrdenCompra"

    headers = [
        "NombreCampo", "TipoDato", "Requerido", "Validacion",
        "Etiqueta", "Pagina", "ModuloDestino", "ValoresEnum", "ValorDefault",
    ]
    ws1.append(headers)

    rows = [
        ["Numero", "AutoNumero", "No", "", "Número de Orden", "Create", "Operaciones", "", ""],
        ["FechaCreacion", "Fecha", "Sí", "", "Fecha de Creación", "Create", "Operaciones", "", ""],
        ["MontoTotal", "Decimal", "Sí", "range:0-999999.99", "Monto Total", "Create", "Operaciones", "", "0.00"],
        ["Descripcion", "Texto", "No", "max_length:500", "Descripción", "Create", "Operaciones", "", ""],
        ["Proveedor", "Texto", "Sí", "min_length:3;max_length:100", "Proveedor", "Create", "Operaciones", "", ""],
        ["Estado", "Enum", "Sí", "", "Estado", "Create", "Operaciones", "Pendiente,Aprobado,Rechazado,Completado", "Pendiente"],
        ["Activo", "Booleano", "No", "", "Activo", "Create", "Operaciones", "", "true"],
        ["CodigoReferencia", "Texto", "No", "regex:^[A-Z]{2}-\\d{4}$;unique", "Código de Referencia", "Create", "Operaciones", "", ""],
    ]
    for row in rows:
        ws1.append(row)

    # ─── Hoja 2: LineaDetalle ───────────────────────────────
    ws2 = wb.create_sheet("LineaDetalle")
    ws2.append(headers)

    rows2 = [
        ["Cantidad", "Entero", "Sí", "range:1-10000", "Cantidad", "Create", "Operaciones", "", "1"],
        ["PrecioUnitario", "Decimal", "Sí", "range:0-99999.99", "Precio Unitario", "", "Operaciones", "", ""],
        ["Subtotal", "Decimal", "No", "", "Subtotal", "", "Operaciones", "", ""],
        ["ProductoNombre", "Texto", "Sí", "max_length:200", "Producto", "", "Operaciones", "", ""],
        ["UnidadMedida", "Enum", "Sí", "", "Unidad de Medida", "", "Operaciones", "Kilogramo,Tonelada,Litro,Unidad", ""],
    ]
    for row in rows2:
        ws2.append(row)

    # ─── Hoja 3: _Metadata (ignorada por el parser) ────────
    ws3 = wb.create_sheet("_Metadata")
    ws3.append(["Clave", "Valor"])
    ws3.append(["Autor", "MendixFormAgent"])
    ws3.append(["Version", "1.0"])

    # Guardar
    output = Path("fixtures/sample_form.xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"✓ Creado: {output}")


if __name__ == "__main__":
    create_sample_form()
