"""Generador de plantilla Excel con headers, dropdowns y ejemplos.

Genera un .xlsx listo para llenar con la definición de entidades.
Incluye:
- Headers de las 15 columnas reconocidas
- Data validation dropdowns para TipoDato, Requerido, VisibleEn, Widget
- Hojas _Relaciones y _Seguridad pre-formateadas
- Filas de ejemplo con una entidad realista
- Anchos de columna configurados

Uso:
    from mendex.templates.excel_template import ExcelTemplateGenerator
    gen = ExcelTemplateGenerator()
    gen.generate(Path("mi_proyecto.xlsx"), module="Operaciones", entity_count=3)

CLI:
    mendex init-excel --module Operaciones --entities 3
"""

from __future__ import annotations

from pathlib import Path

import structlog

logger = structlog.get_logger(__name__)

# Column definitions: (header, width, dropdown_values or None)
ENTITY_COLUMNS: list[tuple[str, int, list[str] | None]] = [
    ("NombreCampo", 20, None),
    ("TipoDato", 14, [
        "Texto", "Entero", "Decimal", "Booleano", "Fecha",
        "FechaHora", "Long", "AutoNumero", "Hash", "Binario",
    ]),
    ("Requerido", 10, ["Si", "No"]),
    ("Validacion", 25, None),
    ("Etiqueta", 20, None),
    ("Pagina", 15, ["Todos", "SoloCrear", "SoloEditar", "SoloOverview"]),
    ("ModuloDestino", 18, None),
    ("ValoresEnum", 30, None),
    ("ValorDefault", 15, None),
    ("Seccion", 15, None),
    ("Widget", 18, [
        "TextBox", "TextArea", "CheckBox", "DatePicker",
        "DropDown", "RadioButtons", "ReferenceSelector",
        "FileManager", "ImageUploader", "RichTextEditor",
    ]),
    ("VisibleEn", 15, ["Todos", "SoloCrear", "SoloEditar", "SoloOverview"]),
    ("Calculado", 25, None),
    ("OrdenSeccion", 12, None),
    ("VisibleSi", 25, None),
]

RELATION_COLUMNS: list[tuple[str, int]] = [
    ("EntidadPadre", 18),
    ("EntidadHija", 18),
    ("TipoRelacion", 14),
    ("Owner", 10),
    ("CascadeDelete", 14),
    ("EsLookup", 10),
    ("Etiqueta", 20),
]

SECURITY_COLUMNS: list[tuple[str, int]] = [
    ("Entidad", 18),
    ("Rol", 18),
    ("PuedeCrear", 12),
    ("PuedeLeer", 12),
    ("PuedeEscribir", 14),
    ("PuedeEliminar", 14),
]

# Example data for a realistic entity
EXAMPLE_ENTITY = "OrdenCompra"
EXAMPLE_ROWS: list[list[str]] = [
    ["Numero", "AutoNumero", "Si", "", "Número de Orden", "Todos", "", "", "", "General", "", "Todos", "", "1", ""],
    ["Fecha", "FechaHora", "Si", "", "Fecha", "Todos", "", "", "", "General", "DatePicker", "Todos", "", "1", ""],
    ["Descripcion", "Texto", "Si", "max_length:500", "Descripción", "Todos", "", "", "", "General", "TextArea", "Todos", "", "1", ""],
    ["MontoTotal", "Decimal", "No", "", "Monto Total", "Todos", "", "", "0.00", "Montos", "", "Todos", "", "2", ""],
    ["Activo", "Booleano", "No", "", "Activo", "SoloOverview", "", "", "true", "Estado", "CheckBox", "Todos", "", "3", ""],
    ["Observaciones", "Texto", "No", "", "Observaciones", "SoloEditar", "", "", "", "Detalle", "RichTextEditor", "SoloEditar", "", "4", ""],
]

EXAMPLE_RELATIONS: list[list[str]] = [
    ["OrdenCompra", "LineaDetalle", "1-*", "default", "Si", "No", "Líneas de detalle"],
    ["OrdenCompra", "TipoOrden", "1-*", "default", "No", "Si", "Tipo de orden (lookup)"],
]

EXAMPLE_SECURITY: list[list[str]] = [
    ["OrdenCompra", "Administrator", "Si", "Si", "Si", "Si"],
    ["OrdenCompra", "User", "Si", "Si", "Si", "No"],
]


class ExcelTemplateGenerator:
    """Genera una plantilla Excel pre-formateada para Mendex."""

    def generate(
        self,
        output_path: Path,
        *,
        module: str = "Operaciones",
        entity_count: int = 1,
        include_examples: bool = True,
    ) -> Path:
        """Genera la plantilla Excel.

        Args:
            output_path: Path de destino para el .xlsx.
            module: Nombre del módulo destino.
            entity_count: Número de hojas de entidad a crear.
            include_examples: Si True, incluye filas de ejemplo.

        Returns:
            Path al archivo generado.
        """
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.datavalidation import DataValidation
        except ImportError:
            raise ImportError(
                "openpyxl es necesario para generar plantillas Excel. "
                "Instálalo con: pip install openpyxl"
            )

        wb = openpyxl.Workbook()

        # Header style
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        # Create entity sheets
        for i in range(entity_count):
            sheet_name = f"Entidad{i + 1}" if entity_count > 1 else EXAMPLE_ENTITY
            if i == 0:
                ws = wb.active
                ws.title = sheet_name
            else:
                ws = wb.create_sheet(sheet_name)

            # Headers
            for col_idx, (header, width, dropdown) in enumerate(ENTITY_COLUMNS, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                ws.column_dimensions[get_column_letter(col_idx)].width = width

                # Add dropdown validation
                if dropdown:
                    dv = DataValidation(
                        type="list",
                        formula1=f'"{",".join(dropdown)}"',
                        allow_blank=True,
                    )
                    dv.error = f"Valor no válido para {header}"
                    dv.errorTitle = "Error de validación"
                    dv.prompt = f"Selecciona un valor para {header}"
                    dv.promptTitle = header
                    col_letter = get_column_letter(col_idx)
                    dv.add(f"{col_letter}2:{col_letter}100")
                    ws.add_data_validation(dv)

            # Set ModuloDestino default
            mod_col = next(
                i for i, (h, _, _) in enumerate(ENTITY_COLUMNS, 1)
                if h == "ModuloDestino"
            )
            for row in range(2, 20):
                ws.cell(row=row, column=mod_col, value=module)

            # Example data
            if include_examples and i == 0:
                for row_idx, row_data in enumerate(EXAMPLE_ROWS, 2):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=value)

        # _Relaciones sheet
        ws_rel = wb.create_sheet("_Relaciones")
        for col_idx, (header, width) in enumerate(RELATION_COLUMNS, 1):
            cell = ws_rel.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws_rel.column_dimensions[get_column_letter(col_idx)].width = width

        # Add relation type dropdown
        dv_rel = DataValidation(
            type="list",
            formula1='"1-*,*-*,1-1"',
            allow_blank=True,
        )
        dv_rel.add("C2:C50")
        ws_rel.add_data_validation(dv_rel)

        # Owner dropdown
        dv_owner = DataValidation(
            type="list",
            formula1='"default,both"',
            allow_blank=True,
        )
        dv_owner.add("D2:D50")
        ws_rel.add_data_validation(dv_owner)

        # Yes/No dropdowns
        for col in ["E", "F"]:
            dv_yn = DataValidation(
                type="list",
                formula1='"Si,No"',
                allow_blank=True,
            )
            dv_yn.add(f"{col}2:{col}50")
            ws_rel.add_data_validation(dv_yn)

        if include_examples:
            for row_idx, row_data in enumerate(EXAMPLE_RELATIONS, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws_rel.cell(row=row_idx, column=col_idx, value=value)

        # _Seguridad sheet
        ws_sec = wb.create_sheet("_Seguridad")
        for col_idx, (header, width) in enumerate(SECURITY_COLUMNS, 1):
            cell = ws_sec.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws_sec.column_dimensions[get_column_letter(col_idx)].width = width

        # Yes/No dropdowns for permission columns
        for col in ["C", "D", "E", "F"]:
            dv_perm = DataValidation(
                type="list",
                formula1='"Si,No"',
                allow_blank=True,
            )
            dv_perm.add(f"{col}2:{col}100")
            ws_sec.add_data_validation(dv_perm)

        if include_examples:
            for row_idx, row_data in enumerate(EXAMPLE_SECURITY, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws_sec.cell(row=row_idx, column=col_idx, value=value)

        # Save
        wb.save(output_path)
        logger.info(
            "excel_template_generated",
            path=str(output_path),
            entities=entity_count,
            module=module,
        )

        return output_path
