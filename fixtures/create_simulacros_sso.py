"""Script para crear el fixture simulacros_sso.xlsx.

Ejecutar: python fixtures/create_simulacros_sso.py

Genera un archivo .xlsx de ejemplo con el formato Excel v2 completo
para el caso GSS-FOR-008 Evaluación de Simulacros.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIXTURES_DIR = Path(__file__).parent

# ─── Estilos ─────────────────────────────────────────────────

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SPECIAL_FILL = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
REQUIRED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _style_headers(ws, fill=None):
    fill = fill or HEADER_FILL
    for cell in ws[1]:
        if cell.value:
            cell.font = HEADER_FONT
            cell.fill = fill
            cell.alignment = ALIGNMENT


def _auto_width(ws):
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)


def create_simulacros_xlsx():
    wb = openpyxl.Workbook()

    # ─── Hoja _Config ────────────────────────────────────────
    ws_config = wb.active
    ws_config.title = "_Config"
    ws_config.append(["Clave", "Valor"])
    ws_config.append(["ModuloDestino", "SSO_Seguridad"])
    ws_config.append(["PrefijoEntidad", "SSO_"])
    ws_config.append(["Layout", "Atlas_Default"])
    ws_config.append(["GenerarMicroflows", "Si"])
    ws_config.append(["GenerarPaginas", "Si"])
    _style_headers(ws_config, SPECIAL_FILL)
    _auto_width(ws_config)

    # ─── Hoja Simulacro ─────────────────────────────────────
    ws_sim = wb.create_sheet("Simulacro")
    headers = [
        "NombreCampo", "TipoDato", "Requerido", "Validacion",
        "Etiqueta", "Pagina", "ModuloDestino", "ValoresEnum", "ValorDefault",
        "Seccion", "Widget", "VisibleEn", "Calculado", "OrdenSeccion",
    ]
    ws_sim.append(headers)
    rows = [
        ["Codigo", "AutoNumero", "", "", "Código", "Overview", "", "", "",
         "Generalidades", "", "SoloOverview", "", "1"],
        ["Sede", "Enum", "Si", "", "Sede", "Create", "", "Lima,Chimbote,Chicama", "",
         "Generalidades", "DropDown", "Todos", "", "1"],
        ["FechaSimulacro", "Fecha", "Si", "", "Fecha del Simulacro", "Create", "", "", "",
         "Generalidades", "", "Todos", "", "1"],
        ["HoraInicio", "Texto", "Si", "", "Hora de Inicio", "Create", "", "", "",
         "Generalidades", "", "Todos", "", "1"],
        ["HoraFin", "Texto", "", "", "Hora de Fin", "Create", "", "", "",
         "Generalidades", "", "Todos", "", "1"],
        ["Lugar", "Texto", "Si", "maximo 200 caracteres", "Lugar", "Create", "", "", "",
         "Generalidades", "", "Todos", "", "1"],
        ["TipoEmergencia", "Enum", "Si", "", "Tipo de Emergencia", "Create", "",
         "Sismo,Incendio,Derrame,Tsunami,Otro", "",
         "Generalidades", "", "Todos", "", "1"],
        ["DescripcionEscenario", "TextoGrande", "", "maximo 500 caracteres",
         "Descripción del Escenario", "Create", "", "", "",
         "Generalidades", "TextArea", "Todos", "", "1"],
        ["Criterio1_Alarma", "Entero", "", "entre 0 y 4",
         "Activación de alarma", "Edit", "", "", "",
         "Evaluacion", "", "SoloEditar", "", "2"],
        ["Criterio2_Evacuacion", "Entero", "", "entre 0 y 4",
         "Evacuación ordenada", "Edit", "", "", "",
         "Evaluacion", "", "SoloEditar", "", "2"],
        ["Criterio3_PuntosReunion", "Entero", "", "entre 0 y 4",
         "Puntos de reunión", "Edit", "", "", "",
         "Evaluacion", "", "SoloEditar", "", "2"],
        ["Criterio4_Brigadistas", "Entero", "", "entre 0 y 4",
         "Respuesta de brigadistas", "Edit", "", "", "",
         "Evaluacion", "", "SoloEditar", "", "2"],
        ["Criterio5_Comunicacion", "Entero", "", "entre 0 y 4",
         "Comunicación durante emergencia", "Edit", "", "", "",
         "Evaluacion", "", "SoloEditar", "", "2"],
        ["PuntajeTotal", "Entero", "", "", "Puntaje Total", "", "", "", "",
         "Evaluacion", "", "Todos",
         "Criterio1_Alarma + Criterio2_Evacuacion + Criterio3_PuntosReunion + Criterio4_Brigadistas + Criterio5_Comunicacion",
         "2"],
        ["Calificacion", "Enum", "", "", "Calificación", "", "",
         "Excelente,Bueno,Regular,Deficiente", "",
         "Evaluacion", "", "Todos", "", "2"],
        ["Observaciones", "TextoGrande", "", "", "Observaciones Generales", "", "", "", "",
         "Observaciones", "TextArea", "Todos", "", "3"],
        ["AccionesCorrectivas", "TextoGrande", "", "", "Acciones Correctivas", "", "", "", "",
         "Observaciones", "TextArea", "SoloEditar", "", "3"],
        ["Estado", "Enum", "Si", "", "Estado", "", "",
         "Borrador,EnRevision,Aprobado,Cerrado", "Borrador",
         "", "", "Todos", "", ""],
        ["FechaAprobacion", "Fecha", "", "", "Fecha de Aprobación", "", "", "", "",
         "", "", "SoloOverview", "", ""],
        ["AprobadoPor", "Texto", "", "", "Aprobado Por", "", "", "", "",
         "", "", "SoloOverview", "", ""],
    ]
    for row in rows:
        ws_sim.append(row)
    _style_headers(ws_sim)
    _auto_width(ws_sim)

    # ─── Hoja Participante ───────────────────────────────────
    ws_part = wb.create_sheet("Participante")
    part_headers = [
        "NombreCampo", "TipoDato", "Requerido", "Validacion",
        "Etiqueta", "Pagina", "ModuloDestino", "ValoresEnum", "ValorDefault",
        "Seccion", "Widget", "VisibleEn",
    ]
    ws_part.append(part_headers)
    part_rows = [
        ["NombreCompleto", "Texto", "Si", "maximo 150 caracteres",
         "Nombre Completo", "Create", "", "", "", "", "", ""],
        ["DNI", "Texto", "Si", "unico",
         "DNI", "Create", "", "", "", "", "", ""],
        ["Area", "Texto", "Si", "",
         "Área", "Create", "", "", "", "", "", ""],
        ["Cargo", "Texto", "", "",
         "Cargo", "", "", "", "", "", "", ""],
        ["Asistio", "Booleano", "", "",
         "Asistió", "Edit", "", "", "false", "", "CheckBox", ""],
        ["RolEnSimulacro", "Enum", "", "",
         "Rol en Simulacro", "", "", "Participante,Brigadista,Observador,Evaluador", "",
         "", "", ""],
        ["Observacion", "TextoGrande", "", "",
         "Observación Individual", "", "", "", "",
         "", "TextArea", "SoloEditar"],
    ]
    for row in part_rows:
        ws_part.append(row)
    _style_headers(ws_part)
    _auto_width(ws_part)

    # ─── Hoja EvidenciaFotografica ───────────────────────────
    ws_evi = wb.create_sheet("EvidenciaFotografica")
    evi_headers = [
        "NombreCampo", "TipoDato", "Requerido", "Etiqueta",
    ]
    ws_evi.append(evi_headers)
    evi_rows = [
        ["Descripcion", "Texto", "Si", "Descripción de la foto"],
        ["Foto", "Imagen", "Si", "Fotografía"],
        ["FechaCaptura", "Fecha", "", "Fecha de captura"],
    ]
    for row in evi_rows:
        ws_evi.append(row)
    _style_headers(ws_evi)
    _auto_width(ws_evi)

    # ─── Hoja _Relaciones ────────────────────────────────────
    ws_rel = wb.create_sheet("_Relaciones")
    ws_rel.append(["EntidadOrigen", "EntidadDestino", "Tipo", "NombreAsociacion", "CascadeDelete"])
    ws_rel.append(["Simulacro", "Participante", "1-*", "Simulacro_Participantes", "Si"])
    ws_rel.append(["Simulacro", "EvidenciaFotografica", "1-*", "Simulacro_Evidencias", "Si"])
    _style_headers(ws_rel, SPECIAL_FILL)
    _auto_width(ws_rel)

    # ─── Hoja _Seguridad ────────────────────────────────────
    ws_seg = wb.create_sheet("_Seguridad")
    ws_seg.append(["Entidad", "Rol", "Crear", "Leer", "Escribir", "Eliminar"])
    ws_seg.append(["*", "Administrator", "Si", "Si", "Si", "Si"])
    ws_seg.append(["Simulacro", "Supervisor_SSO", "Si", "Si", "Si", "No"])
    ws_seg.append(["Simulacro", "Trabajador", "No", "Si", "No", "No"])
    ws_seg.append(["Participante", "Supervisor_SSO", "Si", "Si", "Si", "No"])
    ws_seg.append(["Participante", "Trabajador", "No", "Si", "No", "No"])
    ws_seg.append(["EvidenciaFotografica", "Supervisor_SSO", "Si", "Si", "Si", "Si"])
    ws_seg.append(["EvidenciaFotografica", "Trabajador", "No", "Si", "No", "No"])
    _style_headers(ws_seg, SPECIAL_FILL)
    _auto_width(ws_seg)

    # ─── Guardar ─────────────────────────────────────────────
    output = FIXTURES_DIR / "simulacros_sso.xlsx"
    wb.save(output)
    print(f"Created: {output}")
    return output


if __name__ == "__main__":
    create_simulacros_xlsx()
